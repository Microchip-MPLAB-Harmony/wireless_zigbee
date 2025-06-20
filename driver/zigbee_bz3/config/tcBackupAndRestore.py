from Tools.scripts.ndiff import fopen
import pandas
import openpyxl
from serial import Serial
import serial.tools.list_ports as serialPorts
from enum import Enum
from serial.serialutil import EIGHTBITS, PARITY_NONE, STOPBITS_ONE

# RESTORE = False
RESTORE = True

BAUD_RATE = 115200

# COMMANDS
SOF                 = 36    # $
TC_CMD_BACKUP_REQ   = 1     # 0x0001
TC_CMD_BACKUP_RESP  = 2     # 0x0002
TC_CMD_BACKUP_ACK   = 16    # 0x0010

TC_CMD_RESTORE_REQ  = 3     # 0x0003
TC_CMD_RESTORE_RESP = 4     # 0x0004
TC_CMD_RESTORE_ACK  = 32    # 0x0020

LENGTH_FIELD_SIZE = 2
COMMAND_FIELD_SIZE = 2

total_restore_count = 0
current_restore_index = 0

EXCEL_FILE_NAME = "TC_Backup.xlsx"

class RX_STATE(Enum):
    RX_OFF = 1
    RX_SOF = 2
    RX_LENGTH = 3
    RX_CMD = 4
    RX_DATA = 5

class TC_COMMAND(Enum):
    BACKUP_REQ = 1  # 0x0001
    BACKUP_RESP = 2  # 0x0002
    BACKUP_ACK = 16  # 0x0010
    RESTORE_REQ = 3  # 0x0003
    RESTORE_RESP = 4  # 0x0004
    RESTORE_ACK = 32  # 0x0020


rx_state = RX_STATE.RX_OFF

received_length = 0
received_command = 0
receive_buffer = bytearray()

read_length = 1
transmit_buffer = bytearray()

"""
    summary          : This function will connect to the COM port device
    """
def config_serial():
    global tc_serial
    global serial_com_port
    print("Connecting: " + serial_com_port)
    tc_serial = Serial(
        serial_com_port,
        BAUD_RATE,
        timeout=0.1,
        parity=PARITY_NONE,
        stopbits=STOPBITS_ONE,
        bytesize=EIGHTBITS
    )

"""
    summary          : This function will send an ack in response to the backup request, serially
    """
def send_backup_ack():
    global rx_state
    global transmit_buffer
    global tc_serial
    sof_value = SOF
    length_value = 2
    command_value = TC_CMD_BACKUP_ACK
    transmit_buffer = bytearray()
    transmit_buffer.extend(sof_value.to_bytes(1,"little"))
    transmit_buffer.extend(length_value.to_bytes(2,"little"))
    transmit_buffer.extend(command_value.to_bytes(2, "little"))
    tc_serial.write(transmit_buffer)
    print("\nReq received. Sending ACK")
    rx_state = RX_STATE.RX_SOF

"""
    summary          : This function will store the data received from the com port, to the excel file.
    """
def store_backup_data():
    global receive_buffer
    global received_length
    global tc_backup_sheet
    index = receive_buffer.pop(0)
    line_length = received_length - COMMAND_FIELD_SIZE - 1
    print("Index" + str(index + 1))
    print("Length" + str(line_length))
    print("Store Backup data")
    tc_backup_sheet.cell(row=index+1, column=1, value=hex(line_length))
    for col, val in enumerate(receive_buffer, start=2):
        tc_backup_sheet.cell(row=index+1, column=col, value=hex(val))

"""
    summary          : This function will send a restore request to the device connected to the COM port
    """
def send_restore_request():
    global tc_serial
    global rx_state
    global transmit_buffer
    global current_restore_index
    global total_restore_count
    sof_value = SOF
    length_value = 2
    command_value = TC_CMD_RESTORE_REQ
    transmit_buffer = bytearray()
    transmit_buffer.extend(sof_value.to_bytes(1, "little"))
    transmit_buffer.extend(length_value.to_bytes(2, "little"))
    transmit_buffer.extend(command_value.to_bytes(2, "little"))
    tc_serial.write(transmit_buffer)
    print("\nSending Restore Request")
    rx_state = RX_STATE.RX_SOF

"""
    summary          : This function will read the stored data in excel file and send it serially
    """
def send_restore_data():
    global tc_backup_sheet
    global tc_serial
    global transmit_buffer
    global current_restore_index
    global total_restore_count
    global read_length
    global rx_state
    transmit_buffer = bytearray()
    sof_value = SOF
    command_value = TC_CMD_RESTORE_RESP
    column_count = tc_backup_sheet.max_column
    current_restore_index = current_restore_index+1
    column_length = int(tc_backup_sheet.cell(row=current_restore_index, column=1).value, 16)
    print("\nColumnLength: " + str(column_length))
    length_value = COMMAND_FIELD_SIZE + 1 + column_length
    print("\nLengthValue: " + str(length_value))
    transmit_buffer.extend(sof_value.to_bytes(1, "little"))
    transmit_buffer.extend(length_value.to_bytes(2, "little"))
    transmit_buffer.extend(command_value.to_bytes(2, "little"))
    index_position = current_restore_index - 1
    transmit_buffer.extend(index_position.to_bytes(1, "little"))
    for col_index in range(2, column_length + 2):
        restore_data = int(tc_backup_sheet.cell(row=current_restore_index, column=col_index).value, 16)
        transmit_buffer.append(restore_data)
    print(transmit_buffer)
    tc_serial.write(transmit_buffer)
    rx_state = RX_STATE.RX_SOF

"""
    summary          : This function will evaluate the received command and execute the appropriate action
    """
def process_rx_msg():
    global received_command
    global current_restore_index
    global total_restore_count
    global rx_state
    if received_command == TC_CMD_BACKUP_REQ :
        if RESTORE == True and current_restore_index < total_restore_count:
            print("Restore request")
            send_restore_request()
        else:
            send_backup_ack()
    elif received_command == TC_CMD_BACKUP_RESP :
        store_backup_data()
        send_backup_ack()
    elif received_command == TC_CMD_RESTORE_ACK :
        if current_restore_index < total_restore_count:
            send_restore_data()
        rx_state = RX_STATE.RX_SOF

"""
    summary          : This function will validate the received data and store it in the buffer
    input_parameters : serial_data - the receivced bytes of serail data
    """
def rx_handler(serial_data):
    print(serial_data)
    global read_length
    global rx_state
    global received_length
    global received_command
    global receive_buffer
    match rx_state:
        case RX_STATE.RX_SOF:
            received_data = int.from_bytes(serial_data, byteorder='little')
            if received_data == SOF:
                read_length = LENGTH_FIELD_SIZE
                rx_state = RX_STATE.RX_LENGTH
        case RX_STATE.RX_LENGTH:
            read_length = COMMAND_FIELD_SIZE
            rx_state = RX_STATE.RX_CMD
            received_data = int.from_bytes(serial_data, byteorder='little')
            received_length = received_data
        case RX_STATE.RX_CMD:
            received_data = int.from_bytes(serial_data, byteorder='little')
            received_command = received_data
            read_length = received_length - COMMAND_FIELD_SIZE
            rx_state = RX_STATE.RX_DATA
            if read_length == 0:
                rx_state = RX_STATE.RX_OFF
                read_length = 1
                process_rx_msg()
        case RX_STATE.RX_DATA:
            receive_buffer = bytearray()
            receive_buffer.extend(serial_data)
            rx_state = RX_STATE.RX_OFF
            read_length = 1
            process_rx_msg()

"""
    summary          : This function will read the incoming serial data and send it to rx handler
    """
def read_serial():
    global tc_serial
    global rx_state
    global read_length
    global tc_backup_sheet
    global backup_book
    global total_restore_count
    rx_state = RX_STATE.RX_SOF
    try:
        backup_book = openpyxl.load_workbook(EXCEL_FILE_NAME)
        print("File already exist")
    except FileNotFoundError as e:
        print("File does not exist. Creating...")
        excel_data = pandas.DataFrame()
        excel_data.to_excel(EXCEL_FILE_NAME, index=False)
        backup_book = openpyxl.load_workbook(EXCEL_FILE_NAME)
    tc_backup_sheet = backup_book.active
    print("File Opened")
    if not tc_serial.isOpen():
        print("com cannot not opened")
    else:
        print("com port opened successfully.")
        try:
            if RESTORE == True:
                print("Restore enabled")
                total_restore_count = tc_backup_sheet.max_row
                print("total_restore_count: " + str(total_restore_count))
            else:
                print("Restore disabled")
            while True:
                if tc_serial.in_waiting > 0:
                    rx_data = tc_serial.read(read_length)
                    rx_handler(rx_data)
        except KeyboardInterrupt:
            print("\nInterrupted")
        finally:
            tc_serial.close()
            print("\nConnection closed")
            backup_book.save(EXCEL_FILE_NAME)
            backup_book.close()
            print("\nFile closed")

"""
    summary          : list all the com ports that are connected 
    return_value     : list of COM PORTS
    """
def list_serial_ports():
    ports = serialPorts.comports()
    port_list = []
    print("\nSerial ports list: ")
    for port, desc, hwid in sorted(ports):
        print("{}".format(port))
        port_list.append(str(port))
    return port_list

"""
    summary          : Main function
    """
if __name__ == "__main__":
    global serial_com_port
    print("\nTC backup")

    port_list = list_serial_ports()

    serial_com_port = input("Enter com port: ")
    if serial_com_port in port_list:
        print("com port exist")
        config_serial()
        read_serial()
    else:
        print("Invalid com port")

    print("\nProgram end")

